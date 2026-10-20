from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

SLIDES_DATA = [
    {
        "title": "GlucoTwin AI",
        "subtitle": "Clinical Decision Support & Physiological Digital Twin Platform for Type 2 Diabetes",
        "category": "Unstop Digital Twin Challenge 2026",
        "bullets": [
            "Team: Karpagam Institute of Technology, Coimbatore, Tamil Nadu, India",
            "Domain: Healthcare & AI-driven Clinical Digital Twins",
            "Platform: Real-time CGM Telemetry, XGBoost Spike Forecasts, What-If Simulations"
        ]
    },
    {
        "title": "Executive Summary & Problem Statement",
        "subtitle": "The Challenge in Type 2 Diabetes Management",
        "category": "Clinical Background",
        "bullets": [
            "Clinical Burden: Type 2 Diabetes affects >500M individuals worldwide; postprandial glucose excursions (>10.0 mmol/L) drive vascular complications.",
            "Static EHR Limitation: Conventional clinical records provide only periodic snapshots (HbA1c quarterly) without continuous physiological context.",
            "Wearable Telemetry Disconnect: Modern Continuous Glucose Monitors (CGM) stream massive data, but clinicians lack predictive decision tools.",
            "The Opportunity: A versioned Physiological Digital Twin synthesizing static clinical baselines with real-time biometric telemetry."
        ]
    },
    {
        "title": "Proposed Solution & Healthcare Use Case",
        "subtitle": "Dual-Layer Digital Twin with Forward Risk Forecasting",
        "category": "Solution Architecture",
        "bullets": [
            "Dual-Layer Twin Architecture: Combines static EHR clinical history with a dynamic wearable sliding-window buffer.",
            "2-Hour Spike Forecaster: Machine learning pipeline predicting postprandial hyperglycemia (>= 10.0 mmol/L / 180 mg/dL) 2 hours in advance.",
            "SHAP Feature Attribution: Statistical breakdown explaining exactly which physiological signals (CGM trend, HRV, activity) influence the forecast.",
            "Counterfactual Simulations: Clinicians model hypothetical meals and physical activity before prescribing interventions."
        ]
    },
    {
        "title": "System Architecture & Microservice Topology",
        "subtitle": "Decoupled, Event-Driven Clinical Pipeline",
        "category": "Technical Architecture",
        "bullets": [
            "Presentation Layer: React 18 SPA (Vite, TypeScript, Tailwind CSS dark theme, Recharts telemetry charts).",
            "Application Core: Spring Boot 3.3.0 REST Gateway with Spring Security (ROLE_CLINICIAN) & Flyway database migrations.",
            "Event Streaming Bus: Redis 7 Streams (glucotwin:wearable-events) processing streaming CGM, heart rate, and HRV readings.",
            "Persistence Tier: PostgreSQL 16 storing patient EHR, versioned Digital Twin states, and immutable audit logs.",
            "Machine Learning Service: Python FastAPI container hosting calibrated XGBoost 2.0.3 and SHAP explainability engines."
        ]
    },
    {
        "title": "Physiological Digital Twin State Machine",
        "subtitle": "Dual-Layer Representation & Lifecycle Transitions",
        "category": "Digital Twin Engine",
        "bullets": [
            "Static EHR Layer: Age, biological sex, BMI, diabetes onset year, baseline HbA1c, fasting glucose.",
            "Dynamic Wearable Layer: Rolling CGM glucose window (6-12 points), heart rate, HRV autonomic tone, sleep, and step counts.",
            "Status Lifecycle:",
            "  • INITIALISED: Static EHR loaded; awaiting continuous streaming sensor data.",
            "  • ACTIVE: Live wearable telemetry received within the last 60 minutes.",
            "  • STALE: Telemetry elapsed > 60 minutes; predictions automatically carry data-staleness flags."
        ]
    },
    {
        "title": "AI / Machine Learning Approach & SHAP",
        "subtitle": "Calibrated XGBoost Model with Local Explanations",
        "category": "Machine Learning",
        "bullets": [
            "Algorithm: Pre-trained XGBoost 2.0.3 binary classification model.",
            "Inference Objective: Probability of blood glucose exceeding 10.0 mmol/L within a 2-hour forward horizon.",
            "Outputs Delivered:",
            "  • Point Probability: Calibrated risk probability (0.0% to 100.0%).",
            "  • Categorical Risk Tier: LOW (<30%), MODERATE (30-60%), HIGH (60-85%), CRITICAL (>=85%).",
            "  • 95% Confidence Interval: Empirical bounds evaluating prediction certainty.",
            "  • SHAP TreeExplainer: Directional feature attributions (e.g., cgm_current increases risk by +32.0%)."
        ]
    },
    {
        "title": "What-If Counterfactual Simulations",
        "subtitle": "Hypothetical Scenario Modeling with Transactional Memory Isolation",
        "category": "Clinical Simulation",
        "bullets": [
            "Clinical Purpose: Allows clinicians and researchers to test meal carbohydrate intake (0-500g), physical exercise levels, and medication adherence.",
            "Absolute Memory Isolation Guarantee:",
            "  • Simulation executes exclusively against an in-memory ephemeral twin clone.",
            "  • Observed Digital Twin state and physiological history are strictly IMMUTABLE.",
            "  • Results are tagged with SIMULATED provenance and isolated from real patient EHR.",
            "Verified Invariant: Tests confirm baseline observed glucose remains identical before, during, and after simulation runs."
        ]
    },
    {
        "title": "Multi-Agent Clinical Insights Engine",
        "subtitle": "Collaborative Agent Reasoning with Deterministic Fallbacks",
        "category": "Decision Support",
        "bullets": [
            "RiskEvidenceAgent: Synthesizes EHR baselines, dynamic telemetry, and ML spike probabilities.",
            "ClinicalKnowledgeAgent: Retrieves clinical guidance (ADA Standards of Care, glycemic variability guidelines).",
            "InsightExplanationAgent: Formulates structured medical narrative (LLM-enabled with automatic deterministic rule-based fallback).",
            "InsightVerificationAgent: Verifies numeric claims, asserts safety disclaimers, and records pipeline execution trace."
        ]
    },
    {
        "title": "Testing, Validation & Build Verification",
        "subtitle": "100% Passing Automated Verification Across the Stack",
        "category": "Quality Assurance",
        "bullets": [
            "Frontend Suite: 74 / 74 Vitest component tests passing; strict TypeScript check (0 errors); production build clean (912 modules).",
            "Backend Suite: Spring Boot JUnit 5 and jqwik property tests passing (BUILD SUCCESSFUL); executable bootJar compiled.",
            "ML Service Suite: 29 / 29 Pytest unit tests passing (100% pass rate) covering model deserialization, SHAP attributions, and CI bounds.",
            "End-to-End Workflows: 7 / 7 automated workflow integration scenarios verified (telemetry ingestion, prediction, insights, simulation).",
            "Cohort Persistence: 15 independently versioned patient records verified in PostgreSQL."
        ]
    },
    {
        "title": "Known Limitations & Clinical Disclaimer",
        "subtitle": "Honest Operational Scope & Safety Governance",
        "category": "Governance & Limitations",
        "bullets": [
            "Operational Scope: Designed exclusively for research and clinical decision support; not an FDA/CE-cleared autonomous diagnostic device.",
            "Hardware Emulation: High-throughput telemetry streaming uses Redis event streams rather than physical Bluetooth LE hardware pairings.",
            "Local LLM Mode: Uses deterministic rule-based medical reasoning by default when OPENAI_API_KEY is not configured in local development.",
            "Safety Governance: All screens display explicit clinical safety disclaimers and immutable provenance tags."
        ]
    },
    {
        "title": "Future Scope & Roadmap",
        "subtitle": "Scaling from Decision Support to Closed-Loop Systems",
        "category": "Roadmap",
        "bullets": [
            "Hardware Integration: Direct Bluetooth Low Energy (BLE) drivers for Dexcom G7, Abbott FreeStyle Libre 3, and smart insulin pens.",
            "EHR Interoperability: Production FHIR R4 and HL7 SMART-on-FHIR connectors for hospital Electronic Medical Record systems (Epic, Cerner).",
            "Deep Learning Sequencers: Exploring LSTM and Transformer-based temporal models for continuous multi-step trajectory forecasting.",
            "Automated Insulin Delivery (AID): Exploring research sandbox integrations with open-source artificial pancreas algorithms."
        ]
    }
]

def create_pptx(output_path="docs/submission/GLUCOTWIN_AI_PRESENTATION.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    for slide_data in SLIDES_DATA:
        slide = prs.slides.add_slide(blank_layout)

        # Background
        bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = RGBColor(10, 10, 10)
        bg.line.color.rgb = RGBColor(30, 30, 30)

        # Header Category Pill
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(0.4))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = slide_data["category"].upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = RGBColor(59, 130, 246)

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.5), Inches(0.8))
        tf_t = title_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = slide_data["title"]
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = RGBColor(245, 245, 245)

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(0.5))
        tf_s = sub_box.text_frame
        p_s = tf_s.paragraphs[0]
        p_s.text = slide_data["subtitle"]
        p_s.font.size = Pt(15)
        p_s.font.color.rgb = RGBColor(160, 160, 160)

        # Bullets / Content Card
        card = slide.shapes.add_shape(1, Inches(0.8), Inches(2.2), Inches(11.733), Inches(4.5))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(18, 18, 18)
        card.line.color.rgb = RGBColor(40, 40, 40)

        content_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.4), Inches(11.1), Inches(4.1))
        tf_c = content_box.text_frame
        tf_c.word_wrap = True

        for i, bullet in enumerate(slide_data["bullets"]):
            p = tf_c.add_paragraph() if i > 0 else tf_c.paragraphs[0]
            p.text = f"•  {bullet}"
            p.font.size = Pt(15)
            p.font.color.rgb = RGBColor(220, 220, 220)
            p.space_after = Pt(14)

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
        tf_f = footer_box.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = "GlucoTwin AI — Unstop Digital Twin Challenge 2026 | Karpagam Institute of Technology"
        p_f.font.size = Pt(9)
        p_f.font.color.rgb = RGBColor(100, 100, 100)

    prs.save(output_path)
    print(f"PPTX presentation generated at: {output_path}")

def create_pdf(output_path="docs/submission/GLUCOTWIN_AI_PRESENTATION.pdf"):
    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    cat_style = ParagraphStyle('CatStyle', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9, textColor=colors.HexColor('#2563EB'))
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=22, leading=26, textColor=colors.HexColor('#0F172A'))
    sub_style = ParagraphStyle('SubStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=12, leading=16, textColor=colors.HexColor('#475569'))
    bullet_style = ParagraphStyle('BulletStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=10.5, leading=15, textColor=colors.HexColor('#1E293B'))
    footer_style = ParagraphStyle('FooterStyle', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=10, textColor=colors.HexColor('#94A3B8'))

    story = []

    for i, slide in enumerate(SLIDES_DATA):
        story.append(Paragraph(slide["category"].upper(), cat_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph(slide["title"], title_style))
        story.append(Spacer(1, 4))
        story.append(Paragraph(slide["subtitle"], sub_style))
        story.append(Spacer(1, 14))

        # Bullets inside a clean card table
        bullet_paragraphs = [Paragraph(f"• &nbsp; {b}", bullet_style) for b in slide["bullets"]]
        table_data = [[bp] for bp in bullet_paragraphs]
        t = Table(table_data, colWidths=[720])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
            ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#E2E8F0')),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 7),
            ('TOPPADDING', (0,0), (-1,-1), 7),
            ('LEFTPADDING', (0,0), (-1,-1), 16),
            ('RIGHTPADDING', (0,0), (-1,-1), 16),
        ]))
        story.append(t)
        story.append(Spacer(1, 16))
        story.append(Paragraph("GlucoTwin AI — Unstop Digital Twin Challenge 2026 | Karpagam Institute of Technology", footer_style))

        if i < len(SLIDES_DATA) - 1:
            story.append(PageBreak())

    doc.build(story)
    print(f"PDF presentation generated at: {output_path}")

if __name__ == "__main__":
    create_pptx()
    create_pdf()
